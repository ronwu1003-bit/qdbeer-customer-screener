"""导出器：生成 Excel 待发清单 + CSV 备份"""
import os
from datetime import datetime
from typing import List, Dict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# 邮件模板占位（后续可以扩展）
EMAIL_TEMPLATE_EN = """Subject: {subject}

Dear {contact_name},

We are reaching out from Tsingtao Brewery, China's largest beer exporter, 
to explore potential partnership opportunities. Our premium {product_line} 
products have been well received in {country} and we are looking for 
distributors to expand our presence.

Would you be open to a brief conversation about this?

Best regards,
[Your Name]
[Your Title]
Tsingtao Brewery International Division
"""


def export_to_excel(customers: List[Dict], output_dir: str, filename: str = "") -> str:
    """
    生成带样式的 Excel 待发清单
    返回文件路径
    """
    os.makedirs(output_dir, exist_ok=True)
    if not filename:
        filename = f"qdbeer_customers_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "待发清单"

    # 表头样式
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    whitelist_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    headers = [
        "序号", "公司名称", "官网", "邮箱", "邮箱有效",
        "国家", "白名单", "搜索来源", "匹配关键词", "备注",
    ]

    # 写表头
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

    # 写数据
    for idx, c in enumerate(customers, 1):
        row = idx + 1
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=c.get("company_name", ""))
        ws.cell(row=row, column=3, value=c.get("website", ""))
        ws.cell(row=row, column=4, value=c.get("email", ""))
        # 邮箱有效性显示
        ev = c.get("email_valid", 0)
        ev_text = "✅ 强" if ev >= 2 else ("⚠️ 弱" if ev == 1 else "❌ 无效")
        ws.cell(row=row, column=5, value=ev_text)
        ws.cell(row=row, column=6, value=c.get("country", ""))
        # 白名单高亮
        is_wl = bool(c.get("is_whitelist", 0))
        ws.cell(row=row, column=7, value="⭐ 白名单" if is_wl else "")
        ws.cell(row=row, column=8, value=c.get("search_source", ""))
        ws.cell(row=row, column=9, value=c.get("keywords_matched", ""))
        ws.cell(row=row, column=10, value="")

        # 白名单行高亮
        if is_wl:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).fill = whitelist_fill

    # 调整列宽
    col_widths = [6, 35, 35, 28, 10, 18, 12, 14, 22, 12]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    # 冻结首行
    ws.freeze_panes = "A2"

    wb.save(filepath)
    return filepath


def export_to_csv(customers: List[Dict], output_dir: str, filename: str = "") -> str:
    """生成 CSV 备份"""
    os.makedirs(output_dir, exist_ok=True)
    if not filename:
        filename = f"qdbeer_customers_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    filepath = os.path.join(output_dir, filename)

    df = pd.DataFrame(customers)
    df.to_csv(filepath, index=False, encoding="utf-8-sig")
    return filepath
